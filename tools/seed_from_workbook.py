"""One-time seed: read the through-2025 workbook, add 2026 results, write data/idiot_cup.json.

After this runs once, data/idiot_cup.json is the master record. Annual updates edit
that file (see UPDATING.md), not the workbook.
"""
import json
import pathlib
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
WB = ROOT / "source" / "idiot_cup_v3_thru2025.xlsx"
OUT = ROOT / "data" / "idiot_cup.json"

wb = openpyxl.load_workbook(WB, data_only=True)


def money(v):
    if v in (None, ""):
        return None
    return int(str(v).replace("$", "").replace(",", ""))


def name(v):
    return None if v in (None, "?") else str(v).strip()


# ---------- Master ----------
trips = {}
for r in wb["Master"].iter_rows(min_row=6, values_only=True):
    if not r[1]:
        continue
    y = int(r[1])
    cap1, cap2, capw = name(r[9]), name(r[10]), name(r[11])
    trips[y] = {
        "n": r[0],
        "year": y,
        "dest": r[2],
        "course": r[3],
        "intl": bool(r[4]),
        "golfers": r[5],
        "attendees": [],
        "belt": {"champs": [c for c in (name(r[6]), name(r[7])) if c], "gross": None, "net": None},
        "rc": {"caps": [c for c in (cap1, cap2) if c], "winner": capw, "score": None,
               "teams": None, "matches": None, "pts": None, "side": {}},
        "calcutta": None,
        "ctp": None,
    }
    if r[12]:
        trips[y]["calcutta"] = {"owners": [name(r[12]), name(r[13])], "price": money(r[14]),
                                "pot": money(r[15]), "teams": None}

# ---------- Attendance grid ----------
ws = wb["Attendance"]
years = list(range(2003, 2027))
core, others = [], []
section = core
codes = {}
for r in ws.iter_rows(min_row=5, max_row=43, values_only=True):
    if r[0] and str(r[0]).strip().startswith("Grid only"):
        section = others
        continue
    if not r[0]:
        continue
    p = str(r[0]).strip()
    section.append(p)
    vals = list(r[2:2 + len(years)]) + [None] * len(years)
    for y, v in zip(years, vals):
        if v:
            codes[(p, y)] = str(v)
            trips[y]["attendees"].append(p)
grid_totals = {str(r[0]).strip(): r[1] for r in ws.iter_rows(min_row=5, max_row=43, values_only=True)
               if r[0] and r[1] is not None}

# ---------- RC pts by year ----------
ws = wb["RC Pts by Year"]
rc_years = list(range(2009, 2026))
for r in ws.iter_rows(min_row=5, max_row=32, values_only=True):
    if not r[0]:
        continue
    p = str(r[0]).strip()
    vals = list(r[1:]) + [None] * 60
    for i, y in enumerate(rc_years):
        w, l, team = vals[3 * i], vals[3 * i + 1], vals[3 * i + 2]
        if team in ("W", "L"):
            trips[y]["rc"]["side"][p] = team
        if w is not None and l is not None and (w or l):
            trips[y]["rc"].setdefault("pts", None)
            if trips[y]["rc"]["pts"] is None:
                trips[y]["rc"]["pts"] = {}
            trips[y]["rc"]["pts"][p] = [w, l]

# ---------- Lifetime CTP (no per-year history before 2026) ----------
ctp_base = {}
for r in wb["Lifetime Summary"].iter_rows(min_row=4, max_row=31, values_only=True):
    if r[0] and r[12]:
        ctp_base[str(r[0]).strip()] = r[12]

# ---------- Known belt scores ----------
trips[2003]["belt"]["gross"] = "5 over"  # from prior dashboard
trips[2024]["belt"].update(gross="1 over", net="2 under")
trips[2025]["belt"].update(gross="Even", net="2 under")

# ---------- 2026: Bend, OR (Tetherow), Aug 5-10 ----------
t = trips[2026]
jason = ["Jason L", "Shevin", "Geoff", "Soto", "Eddie C", "EJ", "Geno", "Nardini"]
barrett = ["Barrett", "Lavery", "Gy", "Josh", "Ryan K", "Stephan", "Gil", "Chris D"]
t["dates"] = "Aug 5-10, 2026"
t["dest"] = "Bend, OR"
t["attendees"] = jason + barrett
t["belt"] = {"champs": ["Barrett", "Geoff"], "gross": "Even", "net": "2 under",
             "note": "Shevin/Stephan also shot Even / Net 2 under; Barrett/Geoff won the playoff."}
t["rc"] = {
    "caps": ["Jason L", "Barrett"], "winner": "Jason L", "score": [12.5, 11.5],
    "teams": {"Jason L": jason, "Barrett": barrett},
    "side": {**{p: "W" for p in jason}, **{p: "L" for p in barrett}},
    "matches": [
        {"day": 1, "format": "2-man scramble", "a": ["Jason L", "Shevin"], "b": ["Chris D", "Ryan K"], "pa": 1.75, "pb": 0.25},
        {"day": 1, "format": "2-man scramble", "a": ["EJ", "Geno"], "b": ["Gil", "Gy"], "pa": 0.25, "pb": 1.75},
        {"day": 1, "format": "2-man scramble", "a": ["Soto", "Geoff"], "b": ["Lavery", "Josh"], "pa": 2, "pb": 0},
        {"day": 1, "format": "2-man scramble", "a": ["Eddie C", "Nardini"], "b": ["Barrett", "Stephan"], "pa": 1.5, "pb": 0.5},
        {"day": 2, "format": "2-man best ball", "a": ["Shevin", "Geoff"], "b": ["Gil", "Ryan K"], "pa": 2, "pb": 0},
        {"day": 2, "format": "2-man best ball", "a": ["Eddie C", "Soto"], "b": ["Josh", "Stephan"], "pa": 0.25, "pb": 1.75},
        {"day": 2, "format": "2-man best ball", "a": ["Jason L", "Geno"], "b": ["Chris D", "Lavery"], "pa": 0, "pb": 2},
        {"day": 2, "format": "2-man best ball", "a": ["EJ", "Nardini"], "b": ["Barrett", "Gy"], "pa": 1.75, "pb": 0.25},
        {"day": 3, "format": "Singles", "a": ["Shevin"], "b": ["Josh"], "pa": 1, "pb": 0},
        {"day": 3, "format": "Singles", "a": ["Soto"], "b": ["Stephan"], "pa": 0, "pb": 1},
        {"day": 3, "format": "Singles", "a": ["Nardini"], "b": ["Gil"], "pa": 1, "pb": 0},
        {"day": 3, "format": "Singles", "a": ["Eddie C"], "b": ["Gy"], "pa": 1, "pb": 0},
        {"day": 3, "format": "Singles", "a": ["Geno"], "b": ["Lavery"], "pa": 0, "pb": 1},
        {"day": 3, "format": "Singles", "a": ["Jason L"], "b": ["Barrett"], "pa": 0, "pb": 1},
        {"day": 3, "format": "Singles", "a": ["EJ"], "b": ["Chris D"], "pa": 0, "pb": 1},
        {"day": 3, "format": "Singles", "a": ["Geoff"], "b": ["Ryan K"], "pa": 0, "pb": 1},
    ],
    "pts": None,  # derived from matches below
}
t["calcutta"] = {"owners": ["Geno", "EJ"], "price": 4400, "pot": 23600, "teams": 8}
t["ctp"] = [
    {"day": 1, "wins": {"Stephan": 1.5, "Barrett": 1.5, "Eddie C": 1}},
    {"day": 2, "wins": {"Gy": 1, "Geno": 1, "Gil": 1, "EJ": 1}},
    {"day": 3, "wins": {"Soto": 2, "EJ": 1, "Barrett": 1}},
    {"day": 4, "wins": {"Chris D": 1, "Soto": 1, "Geoff": 1, "Barrett": 1}},
]

pts = {}
for m in t["rc"]["matches"]:
    for p in m["a"]:
        w, l = pts.get(p, [0, 0]); pts[p] = [w + m["pa"], l + m["pb"]]
    for p in m["b"]:
        w, l = pts.get(p, [0, 0]); pts[p] = [w + m["pb"], l + m["pa"]]
t["rc"]["pts"] = pts

# Per-trip Calcutta team counts for earlier years are unknown (avg price shown in workbook only)
for y, tr in trips.items():
    if tr["calcutta"] and tr["calcutta"]["teams"] is None:
        tr["calcutta"]["teams"] = round(tr["golfers"] / 2) if tr["golfers"] else None

data = {
    "meta": {
        "title": "The Idiot Cup",
        "since": 2003,
        "updated": "2026-10-25",
        "minTrips": 3,
        "netFootnote": "Net = adjusted net, off the low team of group",
        "ctpNote": "CTP lifetime totals before 2026 come from the workbook; per-year CTP detail starts in 2026.",
    },
    "players": {"core": core, "others": others, "aka": {"Lavery": "Keith", "Shevin": "Daryl"}},
    "ctpBefore2026": ctp_base,
    "trips": [trips[y] for y in sorted(trips)],
}
OUT.parent.mkdir(exist_ok=True)
OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
print("wrote", OUT)
print("grid totals (workbook):", grid_totals)
